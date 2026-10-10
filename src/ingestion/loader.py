"""Load documents from various formats into plain text + metadata."""
from __future__ import annotations

from pathlib import Path
from typing import Iterator
from dataclasses import dataclass

from pypdf import PdfReader
from docx import Document as DocxDocument


@dataclass
class Document:
    content: str
    source: str
    metadata: dict


def load_pdf(path: Path) -> list[Document]:
    reader = PdfReader(str(path))
    docs = []
    for i, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        if text.strip():
            docs.append(
                Document(
                    content=text.strip(),
                    source=str(path),
                    metadata={"file": path.name, "page": i + 1, "type": "pdf"},
                )
            )
    return docs


def load_docx(path: Path) -> list[Document]:
    doc = DocxDocument(str(path))
    text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    if not text.strip():
        return []
    return [
        Document(
            content=text.strip(),
            source=str(path),
            metadata={"file": path.name, "type": "docx"},
        )
    ]


def load_txt(path: Path) -> list[Document]:
    text = path.read_text(encoding="utf-8", errors="ignore")
    if not text.strip():
        return []
    return [
        Document(
            content=text.strip(),
            source=str(path),
            metadata={"file": path.name, "type": path.suffix.lstrip(".")},
        )
    ]


LOADERS = {
    ".pdf": load_pdf,
    ".docx": load_docx,
    ".txt": load_txt,
    ".md": load_txt,
    ".markdown": load_txt,
}


def load_directory(directory: Path) -> Iterator[Document]:
    """Yield Document objects from all supported files in a directory."""
    if not directory.exists():
        raise FileNotFoundError(f"Directory not found: {directory}")

    for path in sorted(directory.rglob("*")):
        if not path.is_file():
            continue
        loader = LOADERS.get(path.suffix.lower())
        if loader is None:
            continue
        try:
            yield from loader(path)
        except Exception as e:
            print(f"[loader] Failed to load {path}: {e}")
